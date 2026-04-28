from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from pydantic import BaseModel, Field

from ..workspace.permissions import ensure_under_root
from .registry import tool

MAX_INLINE_TEXT_CHARS = 4000


def _require_docx():
    try:
        from docx import Document  # type: ignore
    except Exception as exc:
        raise RuntimeError("python-docx 未安装，请先安装: pip install python-docx") from exc
    return Document


def _resolve_docx_path(workspace_root: Path, relative_path: str) -> Path:
    return ensure_under_root(workspace_root, workspace_root / relative_path)


def _resolve_text_path(workspace_root: Path, relative_path: str) -> Path:
    return ensure_under_root(workspace_root, workspace_root / relative_path)


def _read_text_file(workspace_root: Path, relative_path: str) -> str:
    source = _resolve_text_path(workspace_root, relative_path)
    if not source.exists():
        raise FileNotFoundError(f"文本文件不存在: {source}")
    return source.read_text(encoding="utf-8")


def _validate_inline_text_length(field_name: str, text: str) -> None:
    if len(text) > MAX_INLINE_TEXT_CHARS:
        raise ValueError(
            f"{field_name} 长度为 {len(text)}，超过内联上限 {MAX_INLINE_TEXT_CHARS}。"
            "请先将内容写入文本文件，再使用 *_from_file 工具。"
        )


class DocxCreateArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    title: str | None = Field(default=None, description="可选文档标题")


class DocxAddParagraphArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    text: str = Field(..., description="段落文本")
    style: str | None = Field(default=None, description="段落样式名（可选）")


class DocxAddHeadingArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    text: str = Field(..., description="标题文本")
    level: int = Field(default=1, ge=0, le=9, description="标题级别（0-9）")


class DocxAddTableArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    headers: list[str] = Field(default_factory=list, description="表头")
    rows: list[list[str]] = Field(default_factory=list, description="数据行")


class DocxReplaceTextArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    old_text: str = Field(..., description="待替换文本")
    new_text: str = Field(..., description="替换后文本")


class DocxReadOutlineArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    max_items: int = Field(default=200, ge=1, le=1000, description="最大返回条目数")


class DocxSaveAsArgs(BaseModel):
    source_path: str = Field(..., description="源DOCX路径（相对workspace_root）")
    target_path: str = Field(..., description="目标DOCX路径（相对workspace_root）")


class DocxAddParagraphFromFileArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    text_file_path: str = Field(..., description="文本文件路径（相对workspace_root）")
    style: str | None = Field(default=None, description="段落样式名（可选）")


class DocxAddHeadingFromFileArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    text_file_path: str = Field(..., description="文本文件路径（相对workspace_root）")
    level: int = Field(default=1, ge=0, le=9, description="标题级别（0-9）")


class DocxBuildCheckArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    expected_headings: list[str] = Field(default_factory=list, description="期望包含的标题文本列表")
    min_paragraph_chars: int = Field(default=20, ge=0, le=10_000, description="正文最小字符数阈值")


class DocxInspectStructureArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    max_nodes: int = Field(default=500, ge=1, le=5000, description="最大返回结构节点数量")
    text_preview_chars: int = Field(default=80, ge=10, le=500, description="文本预览截断长度")


class DocxListStylesArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")


class DocxReadTableArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    table_index: int = Field(default=0, ge=0, description="目标表格索引（从0开始）")
    max_rows: int = Field(default=200, ge=1, le=2000, description="最大返回行数")
    max_cols: int = Field(default=50, ge=1, le=200, description="最大返回列数")


class DocxUpdateTableStyleArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    table_index: int = Field(default=0, ge=0, description="目标表格索引（从0开始）")
    style: str | None = Field(default=None, description="表格样式名（如 Table Grid）")
    alignment: str | None = Field(default=None, description="表格对齐方式：left/center/right")
    autofit: bool | None = Field(default=None, description="是否自动适配列宽")


class DocxUpdateCellStyleArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    table_index: int = Field(default=0, ge=0, description="目标表格索引（从0开始）")
    row_start: int = Field(default=0, ge=0, description="起始行（包含，0开始）")
    row_end: int | None = Field(default=None, ge=0, description="结束行（包含，默认最后一行）")
    col_start: int = Field(default=0, ge=0, description="起始列（包含，0开始）")
    col_end: int | None = Field(default=None, ge=0, description="结束列（包含，默认最后一列）")
    bold: bool | None = Field(default=None, description="是否加粗单元格文本")
    italic: bool | None = Field(default=None, description="是否斜体单元格文本")
    text_align: str | None = Field(default=None, description="文本水平对齐：left/center/right/justify")
    vertical_align: str | None = Field(default=None, description="单元格垂直对齐：top/center/bottom")


class DocxReplaceInTableArgs(BaseModel):
    path: str = Field(..., description="DOCX文件路径（相对workspace_root）")
    table_index: int = Field(default=0, ge=0, description="目标表格索引（从0开始）")
    old_text: str = Field(..., description="待替换文本")
    new_text: str = Field(..., description="替换后文本")
    row_start: int = Field(default=0, ge=0, description="起始行（包含，0开始）")
    row_end: int | None = Field(default=None, ge=0, description="结束行（包含，默认最后一行）")
    columns: list[int] | None = Field(default=None, description="仅替换指定列（0开始），默认全部列")
    case_sensitive: bool = Field(default=True, description="是否大小写敏感")


def _table_at(doc: Any, table_index: int) -> Any:
    if table_index >= len(doc.tables):
        raise IndexError(f"表格索引越界: {table_index}，当前仅有 {len(doc.tables)} 个表格")
    return doc.tables[table_index]


def _table_alignment_value(alignment: str | None):
    if alignment is None:
        return None
    normalized = alignment.strip().lower()
    if not normalized:
        return None
    try:
        from docx.enum.table import WD_TABLE_ALIGNMENT  # type: ignore
    except Exception as exc:
        raise RuntimeError("python-docx 未安装，请先安装: pip install python-docx") from exc
    mapping = {
        "left": WD_TABLE_ALIGNMENT.LEFT,
        "center": WD_TABLE_ALIGNMENT.CENTER,
        "right": WD_TABLE_ALIGNMENT.RIGHT,
    }
    if normalized not in mapping:
        raise ValueError("alignment 仅支持 left/center/right")
    return mapping[normalized]


def _normalize_range(start: int, end: int | None, max_index: int, label: str) -> tuple[int, int]:
    if max_index < 0:
        raise ValueError(f"{label} 范围为空")
    resolved_end = max_index if end is None else end
    if start > resolved_end:
        raise ValueError(f"{label} 起始值不能大于结束值: {start} > {resolved_end}")
    if start < 0 or resolved_end > max_index:
        raise IndexError(f"{label} 越界: {start}-{resolved_end}，有效范围 0-{max_index}")
    return start, resolved_end


def _paragraph_alignment_value(value: str | None):
    if value is None:
        return None
    normalized = value.strip().lower()
    if not normalized:
        return None
    try:
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT  # type: ignore
    except Exception as exc:
        raise RuntimeError("python-docx 未安装，请先安装: pip install python-docx") from exc
    mapping = {
        "left": WD_PARAGRAPH_ALIGNMENT.LEFT,
        "center": WD_PARAGRAPH_ALIGNMENT.CENTER,
        "right": WD_PARAGRAPH_ALIGNMENT.RIGHT,
        "justify": WD_PARAGRAPH_ALIGNMENT.JUSTIFY,
    }
    if normalized not in mapping:
        raise ValueError("text_align 仅支持 left/center/right/justify")
    return mapping[normalized]


def _vertical_alignment_value(value: str | None):
    if value is None:
        return None
    normalized = value.strip().lower()
    if not normalized:
        return None
    try:
        from docx.enum.table import WD_ALIGN_VERTICAL  # type: ignore
    except Exception as exc:
        raise RuntimeError("python-docx 未安装，请先安装: pip install python-docx") from exc
    mapping = {
        "top": WD_ALIGN_VERTICAL.TOP,
        "center": WD_ALIGN_VERTICAL.CENTER,
        "bottom": WD_ALIGN_VERTICAL.BOTTOM,
    }
    if normalized not in mapping:
        raise ValueError("vertical_align 仅支持 top/center/bottom")
    return mapping[normalized]


def make_docx_tools(*, workspace_root: Path):
    @tool(name="docx_create", description="创建一个DOCX文档，可选写入标题。")
    def docx_create(args: DocxCreateArgs) -> str:
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        target.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        if args.title:
            doc.add_heading(args.title, level=0)
        doc.save(target)
        return f"已创建DOCX: {target}"

    @tool(name="docx_add_paragraph", description="向DOCX追加段落。")
    def docx_add_paragraph(args: DocxAddParagraphArgs) -> str:
        _validate_inline_text_length("text", args.text)
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        doc = Document(target)
        doc.add_paragraph(args.text, style=args.style)
        doc.save(target)
        return f"已追加段落到: {target}"

    @tool(
        name="docx_add_paragraph_from_file",
        description="从文本文件读取内容并向DOCX追加段落，适合长文案。",
    )
    def docx_add_paragraph_from_file(args: DocxAddParagraphFromFileArgs) -> str:
        text = _read_text_file(workspace_root, args.text_file_path)
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        doc = Document(target)
        doc.add_paragraph(text, style=args.style)
        doc.save(target)
        return f"已从文件追加段落到: {target}"

    @tool(name="docx_add_heading", description="向DOCX追加标题。")
    def docx_add_heading(args: DocxAddHeadingArgs) -> str:
        _validate_inline_text_length("text", args.text)
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        doc = Document(target)
        doc.add_heading(args.text, level=args.level)
        doc.save(target)
        return f"已追加标题到: {target}"

    @tool(
        name="docx_add_heading_from_file",
        description="从文本文件读取内容并向DOCX追加标题，适合长标题素材。",
    )
    def docx_add_heading_from_file(args: DocxAddHeadingFromFileArgs) -> str:
        text = _read_text_file(workspace_root, args.text_file_path)
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        doc = Document(target)
        doc.add_heading(text, level=args.level)
        doc.save(target)
        return f"已从文件追加标题到: {target}"

    @tool(name="docx_add_table", description="向DOCX追加表格。")
    def docx_add_table(args: DocxAddTableArgs) -> str:
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        if not args.headers and not args.rows:
            raise ValueError("headers 和 rows 不能同时为空")

        col_count = len(args.headers) if args.headers else len(args.rows[0])
        if col_count <= 0:
            raise ValueError("无法推断表格列数")

        doc = Document(target)
        table = doc.add_table(rows=1 if args.headers else 0, cols=col_count)
        if args.headers:
            for idx, header in enumerate(args.headers):
                table.rows[0].cells[idx].text = header
        for row_data in args.rows:
            row = table.add_row()
            for idx in range(col_count):
                row.cells[idx].text = row_data[idx] if idx < len(row_data) else ""
        doc.save(target)
        return f"已追加表格到: {target}"

    @tool(name="docx_replace_text", description="替换DOCX段落中的文本。")
    def docx_replace_text(args: DocxReplaceTextArgs) -> str:
        _validate_inline_text_length("old_text", args.old_text)
        _validate_inline_text_length("new_text", args.new_text)
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        doc = Document(target)
        replaced = 0
        for paragraph in doc.paragraphs:
            if args.old_text in paragraph.text:
                paragraph.text = paragraph.text.replace(args.old_text, args.new_text)
                replaced += 1
        doc.save(target)
        return f"已替换 {replaced} 个段落并保存: {target}"

    @tool(name="docx_read_outline", description="读取DOCX大纲（标题和前几段文本）。")
    def docx_read_outline(args: DocxReadOutlineArgs) -> dict[str, object]:
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        doc = Document(target)
        items: list[dict[str, str]] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = str(getattr(para.style, "name", "") or "")
            item_type = "heading" if style_name.lower().startswith("heading") else "paragraph"
            items.append({"type": item_type, "style": style_name, "text": text})
            if len(items) >= args.max_items:
                break
        return {"path": str(target), "items": items, "count": len(items)}

    @tool(name="docx_build_check", description="对 DOCX 构建结果做质量门禁检查（默认仅警告，不阻断）。")
    def docx_build_check(args: DocxBuildCheckArgs) -> dict[str, object]:
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        doc = Document(target)

        heading_texts: list[str] = []
        heading_levels: list[int] = []
        paragraph_texts: list[str] = []
        warnings: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = str(getattr(para.style, "name", "") or "")
            if style_name.lower().startswith("heading"):
                heading_texts.append(text)
                level_match = "".join(ch for ch in style_name if ch.isdigit())
                heading_levels.append(int(level_match) if level_match else 1)
            else:
                paragraph_texts.append(text)
                if args.min_paragraph_chars > 0 and len(text) < args.min_paragraph_chars:
                    warnings.append(f"存在较短正文段落（{len(text)} 字）: {text[:40]}")

        if not any(level == 1 for level in heading_levels):
            warnings.append("缺少一级标题（Heading 1）")
        for i in range(1, len(heading_levels)):
            if heading_levels[i] - heading_levels[i - 1] > 1:
                warnings.append(
                    f"标题层级跳跃: {heading_levels[i - 1]} -> {heading_levels[i]}"
                )
                break

        for expected in args.expected_headings:
            if expected not in heading_texts:
                warnings.append(f"缺少预期标题: {expected}")

        dedup = set(paragraph_texts)
        if paragraph_texts:
            duplicate_ratio = 1.0 - (len(dedup) / len(paragraph_texts))
            if duplicate_ratio >= 0.3:
                warnings.append(f"正文重复比例偏高: {duplicate_ratio:.2%}")
        else:
            duplicate_ratio = 0.0

        return {
            "path": str(target),
            "ok": len(warnings) == 0,
            "warnings": warnings,
            "stats": {
                "heading_count": len(heading_texts),
                "paragraph_count": len(paragraph_texts),
                "duplicate_paragraph_ratio": duplicate_ratio,
            },
        }

    @tool(name="docx_inspect_structure", description="读取DOCX结构（按文档顺序返回段落和表格节点）。")
    def docx_inspect_structure(args: DocxInspectStructureArgs) -> dict[str, object]:
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")

        from docx.table import Table  # type: ignore
        from docx.text.paragraph import Paragraph  # type: ignore

        doc = Document(target)
        nodes: list[dict[str, object]] = []
        paragraph_index = 0
        table_index = 0

        body = doc.element.body
        for child in body.iterchildren():
            if len(nodes) >= args.max_nodes:
                break
            child_tag = str(getattr(child, "tag", ""))
            if child_tag.endswith("}p"):
                para = Paragraph(child, doc)
                text = para.text.strip()
                style_name = str(getattr(para.style, "name", "") or "")
                nodes.append(
                    {
                        "node_id": f"p:{paragraph_index}",
                        "type": "paragraph",
                        "index": paragraph_index,
                        "style": style_name,
                        "text_preview": text[: args.text_preview_chars],
                    }
                )
                paragraph_index += 1
                continue
            if child_tag.endswith("}tbl"):
                table = Table(child, doc)
                row_count = len(table.rows)
                col_count = len(table.columns) if table.rows else 0
                style_name = str(getattr(getattr(table, "style", None), "name", "") or "")
                first_row_preview: list[str] = []
                if table.rows:
                    first_row_preview = [
                        (cell.text or "").strip()[: args.text_preview_chars]
                        for cell in table.rows[0].cells
                    ]
                nodes.append(
                    {
                        "node_id": f"t:{table_index}",
                        "type": "table",
                        "index": table_index,
                        "rows": row_count,
                        "cols": col_count,
                        "style": style_name,
                        "first_row_preview": first_row_preview,
                    }
                )
                table_index += 1

        return {
            "path": str(target),
            "count": len(nodes),
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "nodes": nodes,
        }

    @tool(name="docx_list_styles", description="列出DOCX可用样式（段落/字符/表格）。")
    def docx_list_styles(args: DocxListStylesArgs) -> dict[str, object]:
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")

        from docx.enum.style import WD_STYLE_TYPE  # type: ignore

        doc = Document(target)
        paragraph_styles: list[str] = []
        character_styles: list[str] = []
        table_styles: list[str] = []

        for style in doc.styles:
            style_name = str(getattr(style, "name", "") or "").strip()
            if not style_name:
                continue
            style_type = getattr(style, "type", None)
            if style_type == WD_STYLE_TYPE.PARAGRAPH:
                paragraph_styles.append(style_name)
            elif style_type == WD_STYLE_TYPE.CHARACTER:
                character_styles.append(style_name)
            elif style_type == WD_STYLE_TYPE.TABLE:
                table_styles.append(style_name)

        return {
            "path": str(target),
            "paragraph_styles": sorted(set(paragraph_styles)),
            "character_styles": sorted(set(character_styles)),
            "table_styles": sorted(set(table_styles)),
        }

    @tool(name="docx_read_table", description="读取指定表格内容与元信息。")
    def docx_read_table(args: DocxReadTableArgs) -> dict[str, object]:
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        doc = Document(target)
        table = _table_at(doc, args.table_index)

        rows: list[list[str]] = []
        for row_idx, row in enumerate(table.rows):
            if row_idx >= args.max_rows:
                break
            row_values: list[str] = []
            for col_idx, cell in enumerate(row.cells):
                if col_idx >= args.max_cols:
                    break
                row_values.append((cell.text or "").strip())
            rows.append(row_values)

        style_name = str(getattr(getattr(table, "style", None), "name", "") or "")
        return {
            "path": str(target),
            "table_index": args.table_index,
            "row_count": len(table.rows),
            "col_count": len(table.columns) if table.rows else 0,
            "style": style_name,
            "rows": rows,
            "truncated": len(table.rows) > args.max_rows,
        }

    @tool(name="docx_update_table_style", description="更新指定表格的样式与布局参数。")
    def docx_update_table_style(args: DocxUpdateTableStyleArgs) -> str:
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        doc = Document(target)
        table = _table_at(doc, args.table_index)

        changed: list[str] = []
        if args.style is not None and args.style.strip():
            style_name = args.style.strip()
            try:
                table.style = style_name
            except Exception as exc:
                raise ValueError(f"无法设置表格样式 `{style_name}`，请先用 docx_list_styles 查看可用样式。") from exc
            changed.append(f"style={style_name}")

        alignment_value = _table_alignment_value(args.alignment)
        if alignment_value is not None:
            table.alignment = alignment_value
            changed.append(f"alignment={args.alignment}")

        if args.autofit is not None:
            table.autofit = args.autofit
            changed.append(f"autofit={args.autofit}")

        if not changed:
            raise ValueError("未提供可更新字段，请至少设置 style/alignment/autofit 之一。")

        doc.save(target)
        return f"已更新表格样式: {target} (table_index={args.table_index}, {', '.join(changed)})"

    @tool(name="docx_update_cell_style", description="批量更新指定表格区域的单元格文本/对齐样式。")
    def docx_update_cell_style(args: DocxUpdateCellStyleArgs) -> str:
        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        doc = Document(target)
        table = _table_at(doc, args.table_index)

        row_count = len(table.rows)
        col_count = len(table.columns) if row_count else 0
        if row_count == 0 or col_count == 0:
            raise ValueError("目标表格为空，无法更新单元格样式")

        row_start, row_end = _normalize_range(args.row_start, args.row_end, row_count - 1, "row")
        col_start, col_end = _normalize_range(args.col_start, args.col_end, col_count - 1, "column")

        p_align = _paragraph_alignment_value(args.text_align)
        v_align = _vertical_alignment_value(args.vertical_align)

        if (
            args.bold is None
            and args.italic is None
            and p_align is None
            and v_align is None
        ):
            raise ValueError("未提供可更新字段，请至少设置 bold/italic/text_align/vertical_align 之一。")

        touched = 0
        for r in range(row_start, row_end + 1):
            for c in range(col_start, col_end + 1):
                cell = table.cell(r, c)
                if v_align is not None:
                    cell.vertical_alignment = v_align
                for para in cell.paragraphs:
                    if p_align is not None:
                        para.alignment = p_align
                    for run in para.runs:
                        if args.bold is not None:
                            run.bold = args.bold
                        if args.italic is not None:
                            run.italic = args.italic
                touched += 1

        doc.save(target)
        return (
            f"已更新单元格样式: {target} "
            f"(table_index={args.table_index}, rows={row_start}-{row_end}, cols={col_start}-{col_end}, touched={touched})"
        )

    @tool(name="docx_replace_in_table", description="在指定表格（可限定行列）内替换文本。")
    def docx_replace_in_table(args: DocxReplaceInTableArgs) -> str:
        _validate_inline_text_length("old_text", args.old_text)
        _validate_inline_text_length("new_text", args.new_text)
        if not args.old_text:
            raise ValueError("old_text 不能为空")

        Document = _require_docx()
        target = _resolve_docx_path(workspace_root, args.path)
        if not target.exists():
            raise FileNotFoundError(f"DOCX不存在: {target}")
        doc = Document(target)
        table = _table_at(doc, args.table_index)

        row_count = len(table.rows)
        col_count = len(table.columns) if row_count else 0
        if row_count == 0 or col_count == 0:
            raise ValueError("目标表格为空，无法替换文本")

        row_start, row_end = _normalize_range(args.row_start, args.row_end, row_count - 1, "row")
        if args.columns is None:
            target_cols = list(range(col_count))
        else:
            target_cols = sorted(set(args.columns))
            for col in target_cols:
                if col < 0 or col >= col_count:
                    raise IndexError(f"column 越界: {col}，有效范围 0-{col_count - 1}")

        replace_count = 0
        for r in range(row_start, row_end + 1):
            for c in target_cols:
                cell = table.cell(r, c)
                cell_changed = False
                for para in cell.paragraphs:
                    original = para.text
                    if args.case_sensitive:
                        replaced = original.replace(args.old_text, args.new_text)
                    else:
                        replaced = re.sub(re.escape(args.old_text), args.new_text, original, flags=re.IGNORECASE)
                    if replaced != original:
                        para.text = replaced
                        cell_changed = True
                if cell_changed:
                    replace_count += 1

        doc.save(target)
        return (
            f"已在表格中替换文本: {target} "
            f"(table_index={args.table_index}, rows={row_start}-{row_end}, cols={target_cols}, replaced_cells={replace_count})"
        )

    @tool(name="docx_save_as", description="复制DOCX到新文件。")
    def docx_save_as(args: DocxSaveAsArgs) -> str:
        Document = _require_docx()
        source = _resolve_docx_path(workspace_root, args.source_path)
        target = _resolve_docx_path(workspace_root, args.target_path)
        if not source.exists():
            raise FileNotFoundError(f"源DOCX不存在: {source}")
        target.parent.mkdir(parents=True, exist_ok=True)
        doc = Document(source)
        doc.save(target)
        return f"已另存为: {target}"

    return (
        docx_create,
        docx_add_paragraph,
        docx_add_paragraph_from_file,
        docx_add_heading,
        docx_add_heading_from_file,
        docx_add_table,
        docx_replace_text,
        docx_read_outline,
        docx_build_check,
        docx_inspect_structure,
        docx_list_styles,
        docx_read_table,
        docx_update_table_style,
        docx_update_cell_style,
        docx_replace_in_table,
        docx_save_as,
    )
